#!/usr/bin/env python3
"""Generate the PS-vs-PS UDP performance project proposal deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = Path(__file__).resolve().parent / "PS_UDP_Performance_Comparison.pptx"

SW = 13.333
SH = 7.5
NAVY = "10243A"
NAVY2 = "183650"
TEAL = "11B5A4"
CYAN = "58C8D2"
AMBER = "F2B84B"
CORAL = "F0786B"
GREEN = "47B881"
INK = "1D2B3A"
MUTED = "637487"
PALE = "F3F7FA"
PALE_BLUE = "EAF2F8"
PALE_TEAL = "E7F7F4"
PALE_AMBER = "FFF5E2"
PALE_CORAL = "FFF0EE"
WHITE = "FFFFFF"
LINE = "D8E2EA"
FONT = "Aptos"
MONO = "Consolas"

def color(value):
    return RGBColor.from_string(value)


def set_fill(shape, fill, line=None, width=0.8):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if line:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(width)
    else:
        shape.line.fill.background()


def rect(slide, x, y, w, h, fill, radius=False, line=None, width=0.8):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shape, fill, line, width)
    if radius:
        try:
            shape.adjustments[0] = 0.10
        except (IndexError, TypeError):
            pass
    return shape


def line(slide, x1, y1, x2, y2, stroke=LINE, width=1.2, arrow=False):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color(stroke)
    shape.line.width = Pt(width)
    if arrow:
        shape.line.end_arrowhead = True
    return shape


def text(slide, value, x, y, w, h, size=16, fill=INK, bold=False,
         font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
         margin=0.04, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    for idx, paragraph_text in enumerate(str(value).split("\n")):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.text = paragraph_text
        paragraph.alignment = align
        paragraph.space_after = Pt(3)
        for run in paragraph.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color(fill)
    return box


def rich_text(slide, segments, x, y, w, h, size=16, fill=INK,
              align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(2)
    for label, value, settings in segments:
        run = paragraph.add_run()
        run.text = label
        run.font.name = settings.get("font", FONT)
        run.font.size = Pt(settings.get("size", size))
        run.font.bold = settings.get("bold", False)
        run.font.italic = settings.get("italic", False)
        run.font.color.rgb = color(settings.get("color", fill))
    return box


def bullet_list(slide, items, x, y, w, h, size=15, fill=INK, gap=7,
                bullet_color=TEAL):
    count = max(len(items), 1)
    each_h = h / count
    for idx, item in enumerate(items):
        yy = y + idx * each_h
        rect(slide, x, yy + 0.12, 0.10, 0.10, bullet_color, radius=True)
        text(slide, item, x + 0.20, yy, w - 0.20, each_h - gap / 30,
             size=size, fill=fill)


def card(slide, x, y, w, h, heading, body, accent=TEAL,
         fill=WHITE, heading_size=16, body_size=13):
    rect(slide, x, y, w, h, fill, radius=True, line=LINE)
    rect(slide, x, y, 0.07, h, accent)
    text(slide, heading, x + 0.20, y + 0.16, w - 0.36, 0.35,
         size=heading_size, fill=NAVY, bold=True)
    text(slide, body, x + 0.20, y + 0.56, w - 0.36, h - 0.68,
         size=body_size, fill=MUTED)


def pill(slide, label, x, y, w, fill, foreground=WHITE, size=10):
    rect(slide, x, y, w, 0.30, fill, radius=True)
    text(slide, label, x + 0.05, y + 0.02, w - 0.10, 0.24,
         size=size, fill=foreground, bold=True, align=PP_ALIGN.CENTER,
         valign=MSO_ANCHOR.MIDDLE, margin=0)


def base_slide(prs, title, section, number, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(PALE)
    rect(slide, 0, 0, SW, 0.12, TEAL)
    text(slide, section.upper(), 0.55, 0.30, 4.2, 0.25,
         size=9, fill=TEAL, bold=True)
    text(slide, title, 0.55, 0.64, 12.15, 0.56,
         size=25, fill=NAVY, bold=True)
    if subtitle:
        text(slide, subtitle, 0.58, 1.22, 12.0, 0.38,
             size=12, fill=MUTED)
        content_top = 1.78
    else:
        content_top = 1.48
    line(slide, 0.55, 7.10, 12.78, 7.10, LINE, 0.8)
    text(slide, "PYNQ-Z2  |  PS UDP performance study", 0.58, 7.16,
         7.6, 0.18, size=8, fill=MUTED)
    text(slide, f"{number:02d}", 12.22, 7.14, 0.50, 0.22,
         size=9, fill=MUTED, align=PP_ALIGN.RIGHT)
    return slide, content_top


def table(slide, x, y, widths, row_heights, rows, header_fill=NAVY,
          body_fill=WHITE, font_size=12, first_col_bold=False,
          header_font=11):
    yy = y
    for r_idx, row in enumerate(rows):
        hh = row_heights[r_idx]
        xx = x
        for c_idx, value in enumerate(row):
            ww = widths[c_idx]
            fill = header_fill if r_idx == 0 else (
                PALE_BLUE if r_idx % 2 == 0 else body_fill
            )
            rect(slide, xx, yy, ww, hh, fill, line=WHITE, width=0.7)
            is_header = r_idx == 0
            text(slide, value, xx + 0.09, yy + 0.06, ww - 0.18, hh - 0.10,
                 size=header_font if is_header else font_size,
                 fill=WHITE if is_header else INK,
                 bold=is_header or (first_col_bold and c_idx == 0),
                 valign=MSO_ANCHOR.MIDDLE, margin=0.01)
            xx += ww
        yy += hh


def add_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(NAVY)
    rect(slide, 0, 0, SW, 0.14, TEAL)
    rect(slide, 9.42, 0.14, 3.91, 7.36, NAVY2)
    pill(slide, "PROJECT PROPOSAL  •  MEASUREMENT-DRIVEN", 0.78, 0.86,
         3.60, TEAL, NAVY, 9)
    text(slide, "PS vs PS", 0.78, 1.58, 7.5, 0.90,
         size=42, fill=WHITE, bold=True)
    text(slide, "UDP server performance on the\nPYNQ-Z2", 0.80, 2.60, 8.10, 1.45,
         size=29, fill=WHITE, bold=True)
    text(slide, "Python per-packet baseline  vs  Linux C batched I/O",
         0.83, 4.30, 7.80, 0.50, size=17, fill=CYAN)
    text(slide, "Same ARM processor • same PS GEM • same Ethernet jack",
         0.83, 4.92, 7.80, 0.45, size=13, fill="D9E5EE")
    rect(slide, 0.82, 5.84, 7.95, 0.05, TEAL)
    text(slide, "A fair, reproducible comparison—not a claimed result",
         0.83, 6.07, 7.70, 0.45, size=12, fill="C6D6E3", italic=True)
    # right-side mini architecture illustration
    rect(slide, 9.93, 1.26, 2.76, 1.25, "244760", radius=True)
    text(slide, "PYTHON", 10.17, 1.49, 2.25, 0.30, size=17,
         fill=AMBER, bold=True, align=PP_ALIGN.CENTER)
    text(slide, "recvfrom + sendto", 10.17, 1.91, 2.25, 0.28,
         size=12, fill=WHITE, align=PP_ALIGN.CENTER)
    text(slide, "VS", 10.95, 2.82, 0.75, 0.45, size=21,
         fill=CYAN, bold=True, align=PP_ALIGN.CENTER)
    rect(slide, 9.93, 3.47, 2.76, 1.25, "244760", radius=True)
    text(slide, "C BATCHED", 10.17, 3.70, 2.25, 0.30, size=17,
         fill=TEAL, bold=True, align=PP_ALIGN.CENTER)
    text(slide, "recvmmsg + sendmmsg", 10.07, 4.12, 2.50, 0.28,
         size=11, fill=WHITE, align=PP_ALIGN.CENTER)
    text(slide, "One board. One network path.\nChange only the server strategy.",
         9.75, 5.36, 3.10, 0.74, size=13, fill="D9E5EE",
         align=PP_ALIGN.CENTER)
    text(slide, "PYNQ-Z2 | Zynq-7000 PS | Linux UDP", 0.83, 7.03,
         7.60, 0.22, size=9, fill="A9BDCC")


def add_full_block_diagram(prs, slide_number):
    slide, top = base_slide(
        prs, "Complete proposed-system block diagram", "Appendix | Architecture",
        slide_number,
        "Request path shown left-to-right; the UDP echo response returns over the same Ethernet path."
    )
    # Host outside the board.
    rect(slide, 0.55, top + 1.15, 1.85, 1.45, WHITE, radius=True, line=LINE)
    text(slide, "HOST PC", 0.75, top + 1.36, 1.45, 0.30,
         size=15, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text(slide, "Python benchmark\nWireshark capture", 0.70, top + 1.82,
         1.55, 0.60, size=11, fill=MUTED, align=PP_ALIGN.CENTER)

    # Development board boundary.
    rect(slide, 2.92, top + 0.34, 9.75, 3.52, PALE_BLUE, radius=True,
         line="89AFC8", width=1.5)
    text(slide, "TUL PYNQ-Z2 BOARD", 3.18, top + 0.48, 9.20, 0.28,
         size=13, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)

    blocks = [
        (3.22, top + 1.42, 1.42, 1.13, "RJ45 + PHY",
         "RTL8211E\nbuilt-in jack", TEAL),
        (4.94, top + 1.42, 1.48, 1.13, "PS GEM",
         "Ethernet MAC\nvia MIO", CYAN),
        (6.73, top + 1.42, 1.70, 1.13, "Linux stack",
         "Driver + IP/UDP\nsocket layer", AMBER),
    ]
    for x, y, w, h, heading, body, accent in blocks:
        rect(slide, x, y, w, h, WHITE, radius=True, line=accent, width=1.4)
        text(slide, heading, x + 0.07, y + 0.16, w - 0.14, 0.29,
             size=12, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)
        text(slide, body, x + 0.08, y + 0.54, w - 0.16, 0.48,
             size=10, fill=MUTED, align=PP_ALIGN.CENTER)

    # Two alternative server implementations, run separately.
    rect(slide, 9.40, top + 0.96, 2.94, 0.94, PALE_AMBER, radius=True,
         line="E9C66F", width=1.3)
    text(slide, "A  Python baseline", 9.55, top + 1.11, 2.64, 0.27,
         size=13, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text(slide, "recvfrom() → sendto()", 9.55, top + 1.49, 2.64, 0.23,
         size=10, fill=MUTED, align=PP_ALIGN.CENTER)
    rect(slide, 9.40, top + 2.16, 2.94, 0.94, PALE_TEAL, radius=True,
         line="7CCEC0", width=1.3)
    text(slide, "B  Batched C candidate", 9.55, top + 2.31, 2.64, 0.27,
         size=13, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text(slide, "recvmmsg() → sendmmsg()", 9.55, top + 2.69, 2.64, 0.23,
         size=10, fill=MUTED, align=PP_ALIGN.CENTER)

    # Packet flow and server-selection branch.
    line(slide, 2.40, top + 1.88, 3.18, top + 1.88, TEAL, 2, arrow=True)
    text(slide, "UDP request", 2.35, top + 1.51, 0.88, 0.25,
         size=9, fill=MUTED, align=PP_ALIGN.CENTER)
    line(slide, 4.65, top + 1.98, 4.90, top + 1.98, TEAL, 2, arrow=True)
    line(slide, 6.43, top + 1.98, 6.68, top + 1.98, TEAL, 2, arrow=True)
    line(slide, 8.44, top + 1.98, 9.12, top + 1.98, TEAL, 2, arrow=True)
    line(slide, 9.12, top + 1.45, 9.12, top + 2.63, TEAL, 1.7)
    line(slide, 9.12, top + 1.45, 9.36, top + 1.45, TEAL, 1.7, arrow=True)
    line(slide, 9.12, top + 2.63, 9.36, top + 2.63, TEAL, 1.7, arrow=True)

    rect(slide, 1.10, top + 4.08, 10.95, 0.62, NAVY, radius=True)
    text(slide, "Same PS GEM + Linux stack in both experiments  •  Change only the UDP server implementation",
         1.34, top + 4.23, 10.48, 0.28, size=13, fill=WHITE,
         bold=True, align=PP_ALIGN.CENTER)
    text(slide, "PL and Vivado are not part of this PS-vs-PS network path.",
         1.16, top + 4.82, 10.85, 0.28, size=11, fill=MUTED,
         italic=True, align=PP_ALIGN.CENTER)


def add_source_index(prs, sources, slide_number):
    slide, top = base_slide(
        prs, "Complete source-code appendix", "Appendix | Code", slide_number,
        "The following slides reproduce every current project source file in full; implementation files are split only for readability."
    )
    rows = [["Project source", "Purpose", "Pages"]]
    for rel_path, purpose, _, pages in sources:
        rows.append([rel_path, purpose, str(pages)])
    widths = [4.10, 6.65, 1.10]
    heights = [0.44] + [0.52] * len(sources)
    table(slide, 0.76, top + 0.10, widths, heights, rows, font_size=10,
          first_col_bold=True, header_font=10)
    note_y = top + 0.72 + len(sources) * 0.52
    rect(slide, 0.94, note_y, 11.40, 0.72, PALE_TEAL, radius=True)
    text(slide, "The PPT is a presentation copy. Use the linked repository files as the editable source of truth.",
         1.20, note_y + 0.19, 10.85, 0.34, size=13, fill=NAVY,
         bold=True, align=PP_ALIGN.CENTER)


def add_code_pages(prs, rel_path, purpose, content, first_slide_number,
                   lines_per_slide=22):
    lines = content.splitlines()
    if content.endswith("\n"):
        lines.append("")
    page_count = max(1, (len(lines) + lines_per_slide - 1) // lines_per_slide)
    for page in range(page_count):
        start = page * lines_per_slide
        end = min(start + lines_per_slide, len(lines))
        page_lines = lines[start:end]
        subtitle = (
            f"{purpose}  |  source lines {start + 1}–{end} of {len(lines)}"
        )
        slide, top = base_slide(
            prs, f"Source: {Path(rel_path).name}  •  {page + 1}/{page_count}",
            "Appendix | Full source code",
            first_slide_number + page,
            subtitle,
        )
        rect(slide, 0.66, top + 0.02, 12.02, 5.10, NAVY, radius=True)
        code_box = slide.shapes.add_textbox(
            Inches(0.82), Inches(top + 0.12), Inches(11.72), Inches(4.90)
        )
        frame = code_box.text_frame
        frame.clear()
        frame.word_wrap = False
        frame.margin_left = Inches(0.04)
        frame.margin_right = Inches(0.04)
        frame.margin_top = Inches(0.04)
        frame.margin_bottom = Inches(0.02)
        for line_index, code_line in enumerate(page_lines):
            paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
            paragraph.text = code_line if code_line else " "
            paragraph.space_after = Pt(0)
            paragraph.space_before = Pt(0)
            paragraph.line_spacing = Pt(12.0)
            for run in paragraph.runs:
                run.font.name = MONO
                run.font.size = Pt(9)
                run.font.color.rgb = color("E4EDF4")


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    prs.core_properties.title = "PS vs PS: UDP Server Performance on PYNQ-Z2"
    prs.core_properties.subject = "Literature survey, methodology, tools, components, and proof plan"
    prs.core_properties.author = "Project proposal"
    prs.core_properties.keywords = "PYNQ-Z2, Zynq, UDP, Linux, recvmmsg, sendmmsg, benchmark"
    add_cover(prs)

    # 2: thesis / board reality
    s, top = base_slide(
        prs, "Project in one slide", "01  |  Framing", 2,
        "A practical PS-vs-PS experiment using the PYNQ-Z2 hardware you already have."
    )
    card(s, 0.68, top, 3.78, 2.02, "THE BOARD",
         "The built-in RTL8211E Ethernet jack connects to the Zynq PS GEM through MIO. Linux on the ARM runs the network stack.",
         TEAL, heading_size=16, body_size=13)
    card(s, 4.78, top, 3.78, 2.02, "THE COMPARISON",
         "Baseline: Python recvfrom/sendto per datagram. Candidate: C recvmmsg/sendmmsg, batching up to 32.",
         AMBER, heading_size=16, body_size=13)
    card(s, 8.88, top, 3.78, 2.02, "THE CLAIM",
         "Does batched Linux socket I/O improve measured UDP echo performance on this board, and what happens to loss, latency, and CPU?",
         CYAN, heading_size=16, body_size=13)
    rect(s, 0.68, 4.19, 11.98, 1.40, NAVY, radius=True)
    text(s, "Important boundary", 0.98, 4.43, 2.12, 0.36,
         size=15, fill=CYAN, bold=True)
    text(s, "Both versions run on the PS. The PL is not carrying Ethernet traffic in this selected project.",
         3.15, 4.36, 9.07, 0.54, size=19, fill=WHITE, bold=True,
         valign=MSO_ANCHOR.MIDDLE)
    text(s, "This keeps the comparison testable on the onboard jack and avoids claiming unavailable PL connectivity.",
         0.98, 5.97, 11.4, 0.42, size=13, fill=MUTED)

    # 3: data path
    s, top = base_slide(
        prs, "Where the packets actually go", "01  |  System context", 3,
        "The built-in connector is real hardware; only the server implementation changes."
    )
    steps = [
        ("01", "Host PC", "Python benchmark\ncreates UDP datagrams", NAVY2),
        ("02", "Ethernet link", "RJ45 cable +\nswitch (if used)", TEAL),
        ("03", "PHY + PS GEM", "RTL8211E PHY →\nPS Ethernet MAC", CYAN),
        ("04", "Linux network", "Driver + IP/UDP\nkernel stack", AMBER),
        ("05", "Server process", "Python baseline OR\nC batched candidate", GREEN),
    ]
    x0, boxw, gap = 0.63, 2.30, 0.21
    for idx, (num, heading, body, accent) in enumerate(steps):
        x = x0 + idx * (boxw + gap)
        rect(s, x, top + 0.60, boxw, 1.44, WHITE, radius=True, line=LINE)
        pill(s, num, x + 0.16, top + 0.74, 0.52, accent, NAVY if accent == CYAN else WHITE, 9)
        text(s, heading, x + 0.15, top + 1.14, boxw - 0.30, 0.31,
             size=14, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)
        text(s, body, x + 0.16, top + 1.51, boxw - 0.32, 0.48,
             size=11, fill=MUTED, align=PP_ALIGN.CENTER)
        if idx < len(steps) - 1:
            line(s, x + boxw + 0.015, top + 1.32, x + boxw + gap - 0.03,
                 top + 1.32, TEAL, 1.7, arrow=True)
    rect(s, 1.13, top + 2.62, 11.04, 1.18, PALE_TEAL, radius=True)
    text(s, "Controlled path", 1.43, top + 2.83, 2.2, 0.33,
         size=15, fill=TEAL, bold=True)
    text(s, "Same board • same PS GEM • same Linux kernel • same host test program",
         3.55, top + 2.80, 8.11, 0.42, size=17, fill=NAVY, bold=True,
         valign=MSO_ANCHOR.MIDDLE)
    text(s, "The Ethernet hardware and kernel stack are controls—not the experiment.",
         1.45, top + 3.35, 10.2, 0.30, size=12, fill=MUTED)

    # 4: literature
    s, top = base_slide(
        prs, "Technical literature survey", "02  |  Background", 4,
        "The sources support the design choices; they do not predetermine the result on this board."
    )
    rows = [
        ["Source", "Established point", "Design implication"],
        ["Python socket docs [1]",
         "Python exposes the standard socket API for UDP; it is a usable high-level baseline.",
         "Keep socket semantics and packet bytes the same."],
        ["Linux recvmmsg [2]",
         "One syscall can receive multiple messages; Linux documents potential application performance benefits.",
         "Candidate receive path batches queued datagrams."],
        ["Linux sendmmsg [3]",
         "One syscall can submit multiple messages; partial sends must be handled.",
         "Candidate transmit path retries unsent tail of a batch."],
        ["Linux NAPI docs [4]",
         "NIC/kernel polling and interrupt moderation also affect batching and latency.",
         "Keep kernel/link fixed; application batching is only one layer."],
        ["PYNQ docs [5]",
         "PYNQ supports Python productivity and lower-level C when efficiency matters.",
         "Motivates an embedded Python-versus-C comparison."],
        ["RFC 2544 [6]",
         "Comparable performance reports require defined tests and conditions.",
         "Fix packet sizes, duration, run order, repetitions, and report metrics."],
    ]
    table(s, 0.63, top, [2.25, 5.45, 4.25],
          [0.42, 0.70, 0.70, 0.70, 0.70, 0.70, 0.70],
          rows, font_size=10, first_col_bold=True, header_font=10)
    text(s, "Survey scope: official API / platform documentation + benchmarking guidance (not a claim of exhaustive academic review).",
         0.70, 6.59, 11.7, 0.26, size=10, fill=MUTED, italic=True)

    # 5: gap
    s, top = base_slide(
        prs, "The project gap and research question", "02  |  Background", 5,
        "State the gap narrowly and defensibly: this is a target-board evaluation gap, not a claim that no one has studied batching."
    )
    rect(s, 0.72, top + 0.10, 5.50, 2.00, WHITE, radius=True, line=LINE)
    text(s, "What documentation already gives us", 1.00, top + 0.34, 4.9, 0.34,
         size=16, fill=NAVY, bold=True)
    bullet_list(s, [
        "APIs for ordinary and batched UDP I/O",
        "A plausible mechanism: fewer per-message syscalls",
        "Warnings that kernel/NIC behavior also shapes results",
    ], 1.00, top + 0.83, 4.85, 1.12, size=12, gap=5)
    rect(s, 6.62, top + 0.10, 5.94, 2.00, PALE_AMBER, radius=True, line="F0D89F")
    text(s, "What this project must measure", 6.93, top + 0.34, 5.2, 0.34,
         size=16, fill=NAVY, bold=True)
    bullet_list(s, [
        "Effect on PYNQ-Z2 PS GEM + its installed Linux image",
        "Throughput, loss, latency, and server CPU together",
        "Whether a batched candidate helps under the same offered load",
    ], 6.93, top + 0.83, 5.15, 1.12, size=12, gap=5, bullet_color=AMBER)
    rect(s, 0.72, top + 2.52, 11.84, 1.25, NAVY, radius=True)
    text(s, "Research question", 1.00, top + 2.75, 2.05, 0.33,
         size=14, fill=CYAN, bold=True)
    text(s, "On the PYNQ-Z2, does C batched UDP I/O outperform the Python per-datagram baseline at equal packet sizes and offered load?",
         3.08, top + 2.68, 8.96, 0.62, size=17, fill=WHITE, bold=True,
         valign=MSO_ANCHOR.MIDDLE)
    text(s, "Gap defined by this project: no controlled PYNQ-Z2 measurements are recorded yet. The answer is open until board tests are run.",
         0.89, top + 4.08, 11.5, 0.46, size=12, fill=MUTED, italic=True)
    text(s, "H0: batching does not improve median reply rate.  H1: it improves rate/goodput without increasing loss or corruption.",
         0.89, top + 4.58, 11.5, 0.32, size=10, fill=NAVY, bold=True)

    # 6: proposed components
    s, top = base_slide(
        prs, "Proposed system and components", "03  |  Design", 6,
        "The two implementations share one PS/Linux stack and one benchmark protocol."
    )
    component_cards = [
        ("PYNQ-Z2 board", "Zynq-7000 ARM Cortex-A9 runs Linux; onboard PHY + PS GEM carry packets.", TEAL),
        ("Python baseline", "Blocking UDP socket; receive one datagram and echo it per loop iteration.", AMBER),
        ("C candidate", "Linux recvmmsg/sendmmsg; reusable message arrays; up to 32 datagrams per batch.", GREEN),
        ("Host benchmark", "Sends sequenced payloads in fixed bursts; checks replies and timestamps RTT.", CYAN),
        ("Metrics + report", "JSONL raw runs, median comparison, integrity/loss checks, server CPU evidence.", CORAL),
        ("Network setup", "Wired host-to-board path; same switch/cable, port, negotiated speed, and IP config.", NAVY2),
    ]
    card_w, card_h = 3.75, 1.40
    for idx, (heading, body, accent) in enumerate(component_cards):
        row, col = divmod(idx, 3)
        x = 0.68 + col * 4.04
        y = top + 0.20 + row * 1.70
        card(s, x, y, card_w, card_h, heading, body, accent,
             heading_size=15, body_size=11)
    text(s, "No custom MAC/PHY RTL is required for the selected PS-based experiment.",
         0.80, top + 3.82, 11.7, 0.30, size=12, fill=MUTED, italic=True)

    # 7: two paths
    s, top = base_slide(
        prs, "Baseline and candidate: one controlled A/B test", "03  |  Design", 7,
        "The only intentional independent variable is the server implementation."
    )
    rect(s, 0.76, top + 0.10, 3.25, 3.55, PALE_BLUE, radius=True)
    text(s, "CONTROLLED INPUTS", 1.02, top + 0.34, 2.75, 0.31,
         size=14, fill=NAVY, bold=True)
    bullet_list(s, [
        "Same host benchmark and code",
        "Same UDP packet format",
        "Same payload length and 32-packet burst",
        "Same test duration and socket-buffer request",
        "Same board, Linux, link, and network path",
    ], 1.03, top + 0.86, 2.70, 2.50, size=11, gap=7)
    # central A/B
    rect(s, 4.45, top + 0.15, 3.43, 1.16, PALE_AMBER, radius=True, line="F0D89F")
    text(s, "A  •  Python baseline", 4.72, top + 0.34, 2.9, 0.30,
         size=16, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text(s, "recvfrom → validate → sendto", 4.65, top + 0.75, 3.0, 0.27,
         size=11, fill=MUTED, align=PP_ALIGN.CENTER)
    rect(s, 4.45, top + 2.05, 3.43, 1.16, PALE_TEAL, radius=True, line="A6DDD5")
    text(s, "B  •  C batched candidate", 4.65, top + 2.24, 3.0, 0.30,
         size=16, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text(s, "recvmmsg → validate → sendmmsg", 4.65, top + 2.65, 3.0, 0.27,
         size=11, fill=MUTED, align=PP_ALIGN.CENTER)
    line(s, 6.16, top + 1.33, 6.16, top + 1.94, TEAL, 2, arrow=True)
    rect(s, 8.36, top + 0.10, 4.16, 3.55, NAVY, radius=True)
    text(s, "COMPARE OUTPUTS", 8.68, top + 0.34, 3.5, 0.31,
         size=14, fill=CYAN, bold=True)
    bullet_list(s, [
        "Valid unique replies per second",
        "UDP payload goodput (Mbps)",
        "Loss / corrupt / duplicate replies",
        "Mean, p50 and p95 RTT",
        "Board process CPU time / CPU use",
    ], 8.70, top + 0.86, 3.43, 2.50, size=11, fill=WHITE, gap=7, bullet_color=TEAL)
    text(s, "One server at a time; alternate run order to reduce warm-up / drift bias.",
         1.02, top + 3.98, 11.3, 0.32, size=12, fill=MUTED, italic=True)

    # 8: implementation delta code
    s, top = base_slide(
        prs, "What changes in the code?", "03  |  Design", 8,
        "The experiment targets application-to-kernel call frequency, not the Ethernet PHY or IP stack."
    )
    rect(s, 0.70, top + 0.06, 5.82, 3.44, NAVY, radius=True)
    pill(s, "BASELINE  |  PYTHON", 0.98, top + 0.25, 2.05, AMBER, NAVY, 9)
    code_a = (
        "while running:\n"
        "    payload, peer = sock.recvfrom(1473)\n"
        "    if len(payload) <= 1472:\n"
        "        sock.sendto(payload, peer)"
    )
    text(s, code_a, 1.02, top + 0.78, 5.15, 1.94,
         size=15, fill="DCE8F0", font=MONO)
    text(s, "One receive + one send call per datagram",
         1.02, top + 2.93, 5.12, 0.30, size=12, fill=AMBER, bold=True)
    rect(s, 6.82, top + 0.06, 5.82, 3.44, NAVY, radius=True)
    pill(s, "CANDIDATE  |  LINUX C", 7.10, top + 0.25, 2.30, TEAL, NAVY, 9)
    code_b = (
        "n = recvmmsg(fd, rx, 32, MSG_WAITFORONE, NULL)\n"
        "validate each received datagram\n"
        "sent = 0\n"
        "while sent < valid_count:\n"
        "    sent += sendmmsg(fd, &tx[sent],\n"
        "                     valid_count - sent, 0)"
    )
    text(s, code_b, 7.14, top + 0.75, 5.10, 2.05,
         size=12, fill="DCE8F0", font=MONO)
    text(s, "Up to 32 queued messages per syscall; handle partial sends",
         7.14, top + 2.93, 5.12, 0.30, size=12, fill=TEAL, bold=True)
    rect(s, 0.90, top + 3.83, 11.55, 0.76, PALE_AMBER, radius=True)
    text(s, "Hypothesis: fewer per-datagram user/kernel transitions can reduce CPU cost and raise packet rate when the server is the bottleneck.",
         1.15, top + 4.02, 11.08, 0.34, size=13, fill=NAVY, bold=True,
         align=PP_ALIGN.CENTER)

    # 9: methodology
    s, top = base_slide(
        prs, "Methodology: make the comparison fair", "04  |  Experiment", 9,
        "Pre-register the main metric and keep all other variables fixed."
    )
    phases = [
        ("1", "Inspect", "Record PYNQ image, kernel, CPU architecture, GCC, link speed/duplex, socket buffers."),
        ("2", "Validate", "Run integrity tests against both servers; confirm exact echoes before load testing."),
        ("3", "Warm up", "Warm each server and path equally; do not mix warm-up packets into timed samples."),
        ("4", "Measure", "5+ paired runs per payload; alternate A/B order; save each JSON record."),
        ("5", "Analyze", "Compare medians and variation; examine packet loss, RTT, CPU, and errors."),
        ("6", "Report", "Keep raw logs, environment, commands, captures, plots, and limitations."),
    ]
    for idx, (num, heading, body) in enumerate(phases):
        col, row = idx % 3, idx // 3
        x = 0.74 + col * 4.15
        y = top + 0.12 + row * 1.88
        rect(s, x, y, 3.82, 1.55, WHITE, radius=True, line=LINE)
        pill(s, num, x + 0.18, y + 0.19, 0.40, TEAL, NAVY, 10)
        text(s, heading, x + 0.72, y + 0.20, 2.78, 0.30,
             size=15, fill=NAVY, bold=True)
        text(s, body, x + 0.22, y + 0.65, 3.36, 0.72,
             size=11, fill=MUTED)
    rect(s, 0.84, top + 4.02, 11.68, 0.64, PALE_TEAL, radius=True)
    text(s, "Repeat for 64 B, 512 B and 1472 B payloads; same 32-packet host burst and 10 s offered-load window.",
         1.10, top + 4.19, 11.15, 0.30, size=13, fill=NAVY,
         bold=True, align=PP_ALIGN.CENTER)

    # 10: metrics
    s, top = base_slide(
        prs, "What exactly will be measured?", "04  |  Experiment", 10,
        "Application-level metrics are not a substitute for PHY line-rate testing."
    )
    metric_cards = [
        ("GOODPUT", "Unique echoed payload bits / timed seconds", "Mbps = unique_replies × payload_bytes × 8 / seconds / 10⁶", TEAL),
        ("PACKET RATE", "Count only valid, unique echoes", "replies/s = unique_replies / timed seconds", CYAN),
        ("LOSS", "Sent sequences with no valid echo after drain", "loss % = (sent − unique_replies) / sent × 100", CORAL),
        ("LATENCY", "Host monotonic send-to-valid-echo time", "Report mean, p50, p95; compare under same offered load", AMBER),
        ("CPU COST", "Server process CPU time + Linux CPU monitor", "Report CPU seconds and average utilization separately", GREEN),
        ("INTEGRITY", "Sequence, exact length and payload bytes", "Any corruption / unexpected duplicate is a correctness failure", NAVY2),
    ]
    for idx, (heading, body, measure, accent) in enumerate(metric_cards):
        col, row = idx % 3, idx // 3
        x = 0.70 + col * 4.02
        y = top + 0.12 + row * 1.83
        rect(s, x, y, 3.70, 1.58, WHITE, radius=True, line=LINE)
        text(s, heading, x + 0.18, y + 0.15, 3.32, 0.25,
             size=12, fill=accent, bold=True)
        text(s, body, x + 0.18, y + 0.48, 3.32, 0.43,
             size=12, fill=NAVY, bold=True)
        text(s, measure, x + 0.18, y + 1.01, 3.32, 0.45,
             size=9, fill=MUTED)
    text(s, "Use the same timer boundary and drain interval for A and B; disclose that loopback throughput includes host and Linux overhead.",
         0.82, top + 3.97, 11.50, 0.42, size=11, fill=MUTED, italic=True)

    # 11: test matrix
    s, top = base_slide(
        prs, "Test matrix and controlled variables", "04  |  Experiment", 11,
        "A small matrix is easier to repeat, audit, and explain in a viva."
    )
    matrix_rows = [
        ["Variable", "Plan", "Reason"],
        ["Server", "Python baseline vs batched C", "Independent variable"],
        ["Payload size", "64 B / 512 B / 1472 B", "Small, mid-size, near-MTU behavior"],
        ["Host burst", "32 datagrams (both versions)", "Creates queue opportunities for batch receive"],
        ["Timed load", "10 s + identical 1 s reply drain", "Common throughput window and loss accounting"],
        ["Socket buffers", "Request 4 MiB; record actual values", "Avoid hidden A/B buffer mismatch"],
        ["Repetitions", "At least 5 paired runs per size", "Expose run-to-run variability"],
        ["Order", "Alternate Python→C and C→Python", "Reduce warm-up / thermal / time-order bias"],
        ["Environment", "Same wired host, board image, link, port, background load", "Hold non-treatment variables constant"],
    ]
    table(s, 0.75, top + 0.05, [2.13, 4.72, 4.64],
          [0.42] + [0.47] * 8, matrix_rows, font_size=10, first_col_bold=True,
          header_font=10)
    rect(s, 0.90, top + 4.24, 11.48, 0.55, PALE_CORAL, radius=True)
    text(s, "Record negotiated Ethernet rate; do not call application goodput “1 Gb/s line rate.”",
         1.12, top + 4.38, 11.02, 0.25, size=12, fill=NAVY, bold=True,
         align=PP_ALIGN.CENTER)

    # 12: candidate benefits/hypotheses
    s, top = base_slide(
        prs, "Expected change vs verified improvement", "05  |  Evaluation", 12,
        "Separate the mechanism you implemented from what the data proves."
    )
    rows = [
        ["Implementation delta", "Potential benefit (hypothesis)", "Evidence required"],
        ["recvmmsg() receives multiple queued UDP datagrams in one call",
         "Fewer receive syscalls per packet; lower server CPU per datagram",
         "C process CPU + packets/s under matched traffic"],
        ["sendmmsg() sends multiple datagrams in one call",
         "Fewer transmit syscalls; higher reply rate at high packet rate",
         "Median goodput/replies/s; verify partial-send handling"],
        ["Reusable arrays and buffers",
         "Less per-packet setup/allocation work in the loop",
         "Code review + CPU / rate data; no payload corruption"],
        ["One measurement harness for A and B",
         "Results become directly comparable and reproducible",
         "Same JSON fields, sequence checks, config and test schedule"],
    ]
    table(s, 0.70, top + 0.12, [3.35, 4.25, 3.88],
          [0.48, 0.91, 0.91, 0.91, 0.91], rows,
          font_size=11, first_col_bold=True, header_font=10)
    rect(s, 0.86, top + 4.26, 11.60, 0.62, NAVY, radius=True)
    text(s, "Candidate ≠ proven faster. Only board measurements can support an improvement claim.",
         1.14, top + 4.42, 11.05, 0.28, size=15, fill=WHITE, bold=True,
         align=PP_ALIGN.CENTER)

    # 13: proof / decision
    s, top = base_slide(
        prs, "How to prove the candidate is better", "05  |  Evaluation", 13,
        "Define success before looking at the results."
    )
    gates = [
        ("1", "Correctness gate", "100% exact payload echoes in the integrity test; zero corrupt replies in benchmark."),
        ("2", "Performance gate", "Median goodput or valid replies/s improves consistently across paired trials."),
        ("3", "Reliability gate", "No unacceptable loss increase; disclose loss and all excluded/failed runs."),
        ("4", "Trade-off gate", "Show RTT and CPU alongside throughput—do not hide regressions."),
    ]
    for i, (num, heading, body) in enumerate(gates):
        y = top + 0.10 + i * 0.95
        rect(s, 0.82, y, 7.35, 0.78, WHITE, radius=True, line=LINE)
        pill(s, num, 1.02, y + 0.20, 0.40, TEAL if i < 2 else NAVY2, WHITE, 10)
        text(s, heading, 1.62, y + 0.14, 2.14, 0.27,
             size=13, fill=NAVY, bold=True)
        text(s, body, 3.78, y + 0.12, 4.12, 0.52,
             size=10, fill=MUTED, valign=MSO_ANCHOR.MIDDLE)
    rect(s, 8.56, top + 0.10, 3.82, 3.63, PALE_TEAL, radius=True)
    text(s, "Report the paired change", 8.86, top + 0.38, 3.25, 0.36,
         size=15, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text(s, "Improvement % =\n(median_C − median_Python)\n÷ median_Python × 100",
         8.91, top + 1.08, 3.12, 1.23, size=15, fill=TEAL, bold=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(s, "For loss, report percentage-point change. Show all repeated values or a box plot—not only the best run.",
         8.90, top + 2.60, 3.10, 0.82, size=11, fill=MUTED,
         align=PP_ALIGN.CENTER)
    text(s, "If results are noisy or overlap, conclude “no clear improvement in this setup” and discuss why.",
         0.98, top + 4.13, 11.25, 0.38, size=13, fill=CORAL, bold=True,
         align=PP_ALIGN.CENTER)

    # 14: tools
    s, top = base_slide(
        prs, "Tools and what each one proves", "06  |  Execution", 14,
        "This PS-based project does not need Vivado for its Ethernet data path."
    )
    tool_rows = [
        ["Tool", "Use", "Evidence produced"],
        ["PYNQ-Z2 + Linux", "Run both servers on ARM/PS GEM", "Board-side behavior and CPU environment"],
        ["GCC on PYNQ", "Build C candidate using documented flags", "Compiler and target environment are recorded"],
        ["Host Python 3", "Loopback + same load generator for both servers", "Integrity, sequence counts, RTT and rate"],
        ["compare_results.py", "Compute median metrics from JSONL files", "Reproducible A/B summary"],
        ["top / pidstat", "Observe server CPU under load", "CPU-use trade-off (if available)"],
        ["Wireshark", "Capture host Ethernet interface", "Check ARP, UDP source/destination and packet sizes"],
        ["Git + results/", "Version code, raw logs, configs and report", "Reproducibility and review trail"],
    ]
    table(s, 0.78, top + 0.10, [2.45, 5.10, 3.93],
          [0.43] + [0.59] * 7, tool_rows, font_size=10, first_col_bold=True,
          header_font=10)
    text(s, "Vivado/JTAG is not in scope: it cannot route the onboard PS Ethernet jack into custom PL logic.",
         0.86, top + 4.56, 11.48, 0.35, size=12, fill=MUTED, italic=True)

    # 15: deliverables / schedule
    s, top = base_slide(
        prs, "Execution plan and deliverables", "06  |  Execution", 15,
        "A sequence that gets a defensible result before presentation day."
    )
    plan = [
        ("A", "Environment check", "Boot/network, identify IP and link, verify GCC + Linux APIs.", "Log system versions"),
        ("B", "Build + functional test", "Compile C server; run exact-echo test against A and B.", "Save commands/logs"),
        ("C", "Benchmark", "3 payload sizes × 5+ repeats × alternating A/B order.", "Save JSONL files"),
        ("D", "Analyze", "Compare medians, run-to-run spread, loss, RTT and CPU.", "Create figures/table"),
        ("E", "Conclude", "Explain whether batching helped on this setup and limitations.", "Complete results report"),
    ]
    for idx, (mark, heading, body, proof) in enumerate(plan):
        y = top + 0.12 + idx * 0.86
        accent = [TEAL, CYAN, AMBER, GREEN, CORAL][idx]
        rect(s, 0.82, y, 11.62, 0.69, WHITE, radius=True, line=LINE)
        pill(s, mark, 1.02, y + 0.17, 0.42, accent, NAVY, 10)
        text(s, heading, 1.70, y + 0.11, 2.12, 0.27,
             size=13, fill=NAVY, bold=True)
        text(s, body, 3.82, y + 0.10, 5.55, 0.45,
             size=10, fill=MUTED, valign=MSO_ANCHOR.MIDDLE)
        text(s, proof, 9.45, y + 0.14, 2.64, 0.35,
             size=10, fill=accent, bold=True, align=PP_ALIGN.RIGHT,
             valign=MSO_ANCHOR.MIDDLE)

    # 16: results placeholder, clearly not actual
    s, top = base_slide(
        prs, "Results slide: fill only after board testing", "07  |  Results", 16,
        "No board measurements are available yet. This slide is a reporting template, not a result."
    )
    pill(s, "DATA PENDING", 10.77, 0.37, 1.45, CORAL, WHITE, 9)
    result_rows = [
        ["Payload", "Variant", "Median replies/s", "Goodput Mbps", "Loss %", "p95 RTT ms", "CPU %"],
        ["64 B", "Python baseline", "—", "—", "—", "—", "—"],
        ["64 B", "Batched C", "—", "—", "—", "—", "—"],
        ["512 B", "Python baseline", "—", "—", "—", "—", "—"],
        ["512 B", "Batched C", "—", "—", "—", "—", "—"],
        ["1472 B", "Python baseline", "—", "—", "—", "—", "—"],
        ["1472 B", "Batched C", "—", "—", "—", "—", "—"],
    ]
    table(s, 0.70, top + 0.13, [1.20, 2.10, 2.12, 1.75, 1.35, 1.65, 1.10],
          [0.48] + [0.55] * 6, result_rows, font_size=11,
          first_col_bold=True, header_font=9)
    rect(s, 0.91, top + 4.20, 11.40, 0.66, PALE_AMBER, radius=True)
    text(s, "Do not report localhost smoke-test values as PYNQ-Z2 performance.",
         1.15, top + 4.37, 10.92, 0.30, size=14, fill=NAVY, bold=True,
         align=PP_ALIGN.CENTER)

    # 17: limitations / defense
    s, top = base_slide(
        prs, "Limitations to acknowledge in the viva", "07  |  Results", 17,
        "Clear limitations strengthen the credibility of the conclusion."
    )
    limitation_cards = [
        ("Not raw Ethernet capacity", "Application goodput includes host, kernel, driver and socket costs.", CORAL),
        ("Linux version matters", "recvmmsg/sendmmsg availability and implementation depend on the board image/kernel.", AMBER),
        ("Host can bottleneck", "A saturated generator or receiver can hide server differences.", CYAN),
        ("Burst size changes load", "Fixed host burst may favor batching; disclose and test additional burst settings if time permits.", TEAL),
        ("One board is one sample", "Conclusions apply to this hardware/software configuration, not every Linux system.", GREEN),
        ("CPU metrics differ", "Process CPU time and system-wide utilization are different; label the measurement method.", NAVY2),
    ]
    for idx, (heading, body, accent) in enumerate(limitation_cards):
        col, row = idx % 3, idx // 3
        x = 0.70 + col * 4.02
        y = top + 0.13 + row * 1.77
        card(s, x, y, 3.69, 1.52, heading, body, accent,
             heading_size=13, body_size=10)
    text(s, "A useful result can be “no measurable benefit under this offered load”—that still answers the research question.",
         0.86, top + 3.94, 11.55, 0.47, size=13, fill=NAVY, bold=True,
         align=PP_ALIGN.CENTER)

    # 18: take-away + citations
    s, top = base_slide(
        prs, "Take-away and references", "08  |  Close", 18,
        "What makes the project stronger is the experimental design and honest evidence—not a promised speedup."
    )
    rect(s, 0.74, top + 0.10, 4.35, 3.92, NAVY, radius=True)
    text(s, "THE PLAN", 1.04, top + 0.39, 3.6, 0.30,
         size=13, fill=CYAN, bold=True)
    text(s, "Keep the board and network path fixed.\n\nChange Python per-packet I/O to Linux C batched I/O.\n\nProve correctness first; then compare rate, loss, latency and CPU over repeated runs.",
         1.04, top + 0.90, 3.72, 2.55, size=16, fill=WHITE, bold=True)
    text(s, "No improvement is claimed until board data supports it.",
         1.04, top + 3.54, 3.70, 0.30, size=11, fill=AMBER, bold=True)
    text(s, "REFERENCES", 5.43, top + 0.16, 6.65, 0.30,
         size=13, fill=TEAL, bold=True)
    refs = [
        "[1] Python socket API — docs.python.org/3/library/socket.html",
        "[2] Linux recvmmsg(2) — man7.org/linux/man-pages/man2/recvmmsg.2.html",
        "[3] Linux sendmmsg(2) — man7.org/linux/man-pages/man2/sendmmsg.2.html",
        "[4] Linux NAPI — docs.kernel.org/networking/napi.html",
        "[5] PYNQ introduction — pynq.readthedocs.io/en/latest/",
        "[6] RFC 2544 — rfc-editor.org/rfc/rfc2544 (methodology inspiration; this is not RFC2544 compliance)",
        "[7] AMD UG585 — docs.amd.com/v/u/en-US/ug585-Zynq-7000-TRM",
    ]
    bullet_list(s, refs, 5.46, top + 0.65, 6.67, 3.25,
                size=9, fill=INK, gap=5, bullet_color=TEAL)
    text(s, "Sources accessed 7 Oct 2026. Platform-specific facts and all performance values should be verified against the actual board image and measurements.",
         5.49, top + 4.13, 6.65, 0.45, size=9, fill=MUTED, italic=True)

    project_root = Path(__file__).resolve().parents[1]
    source_specs = [
        ("sw/udp_echo_server.py",
         "Python baseline: one recvfrom/sendto per datagram."),
        ("sw/udp_echo_server_batch.c",
         "C candidate: Linux recvmmsg/sendmmsg batched echo server."),
        ("sw/throughput_benchmark.py",
         "Host benchmark: fixed bursts, sequence validation and metrics."),
        ("sw/udp_loopback_test.py",
         "Host functional test: exact payload echo verification."),
        ("sw/compare_results.py",
         "Analysis tool: compare repeated-run medians."),
    ]
    sources = []
    for relative_path, purpose in source_specs:
        source_text = (project_root / relative_path).read_text(encoding="utf-8")
        source_lines = source_text.splitlines()
        if source_text.endswith("\n"):
            source_lines.append("")
        page_count = max(1, (len(source_lines) + 21) // 22)
        sources.append((relative_path, purpose, source_text, page_count))

    add_full_block_diagram(prs, len(prs.slides) + 1)
    add_source_index(prs, sources, len(prs.slides) + 1)
    next_slide = len(prs.slides) + 1
    for relative_path, purpose, source_text, page_count in sources:
        add_code_pages(prs, relative_path, purpose, source_text, next_slide)
        next_slide += page_count

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(f"Created {OUTPUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build_deck()
